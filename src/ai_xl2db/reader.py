"""Read spreadsheets (.xlsx, .csv) into sheets with a header row and data rows."""

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".xlsx", ".csv"}

_INT = re.compile(r"^[+-]?\d+$")
_FLOAT = re.compile(r"^[+-]?(\d+\.\d*|\.\d+|\d+)([eE][+-]?\d+)?$")


@dataclass
class Sheet:
    title: str
    header: list[object]
    rows: list[list[object]]
    # (row index in rows, column index) -> (formula text, value stored in the file)
    formulas: dict[tuple[int, int], tuple[str, object]] = field(default_factory=dict)
    row_numbers: list[int] = field(default_factory=list)  # spreadsheet row number of each row, 1-based


def _is_empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _find_header(grid: list[list[object]]) -> int:
    """Index of the header row: the first row with at least two filled cells.

    Single-cell rows above it are treated as titles. If no row has two filled
    cells, the first filled row is the header.
    """
    for index, row in enumerate(grid):
        if sum(not _is_empty(v) for v in row) >= 2:
            return index
    for index, row in enumerate(grid):
        if any(not _is_empty(v) for v in row):
            return index
    return -1


def sheet_from_grid(
    title: str, grid: list[list[object]], formulas: dict[tuple[int, int], tuple[str, object]] | None = None
) -> Sheet | None:
    """Build a Sheet from a grid of cell values (row-major, 0-based)."""
    header_index = _find_header(grid)
    if header_index < 0:
        return None

    header_row = grid[header_index]
    width = max(
        (i + 1 for i, v in enumerate(header_row) if not _is_empty(v)),
        default=0,
    )
    if width == 0:
        return None

    header = list(header_row[:width])
    rows = []
    source_rows = []
    for index in range(header_index + 1, len(grid)):
        row = list(grid[index][:width]) + [None] * max(0, width - len(grid[index]))
        if all(_is_empty(v) for v in row):
            continue
        rows.append(row)
        source_rows.append(index)

    # Re-key formulas from grid positions to positions within rows.
    row_position = {source: pos for pos, source in enumerate(source_rows)}
    kept_formulas: dict[tuple[int, int], tuple[str, object]] = {}
    for (grid_row, col), text in (formulas or {}).items():
        if grid_row in row_position and col < width:
            kept_formulas[(row_position[grid_row], col)] = text

    return Sheet(
        title=title,
        header=header,
        rows=rows,
        formulas=kept_formulas,
        row_numbers=[index + 1 for index in source_rows],
    )


def _parse_csv_cell(text: str) -> object:
    """Read a CSV cell as a number when that is unambiguous.

    Values with leading zeros (postcodes, IDs) stay as text so that no digits are lost.
    """
    if text.strip() == "":
        return None
    stripped = text.strip()
    if len(stripped) > 1 and stripped.lstrip("+-")[:1] == "0" and not stripped.lstrip("+-").startswith("0."):
        return text
    if _INT.match(stripped):
        return int(stripped)
    if _FLOAT.match(stripped):
        return float(stripped)
    return text


def read_csv(path: Path) -> list[Sheet]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        text = path.read_text(encoding="gb18030")  # common for Chinese-language CSV exports
    grid = [[_parse_csv_cell(cell) for cell in row] for row in csv.reader(text.splitlines())]
    sheet = sheet_from_grid(path.stem, grid)
    return [sheet] if sheet else []


def read_xlsx(path: Path) -> list[Sheet]:
    import openpyxl

    values_book = openpyxl.load_workbook(path, data_only=True)
    formula_book = openpyxl.load_workbook(path, data_only=False)
    sheets = []
    for ws in values_book.worksheets:
        grid = [list(row) for row in ws.iter_rows(values_only=True)]
        formulas = {}
        formula_ws = formula_book[ws.title]
        for row_index, row in enumerate(formula_ws.iter_rows(values_only=True)):
            for col_index, cell in enumerate(row):
                if isinstance(cell, str) and cell.startswith("="):
                    cached = grid[row_index][col_index] if col_index < len(grid[row_index]) else None
                    formulas[(row_index, col_index)] = (cell, cached)
        sheet = sheet_from_grid(ws.title, grid, formulas)
        if sheet:
            sheets.append(sheet)
    return sheets


def read_spreadsheet(path: Path) -> list[Sheet]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    if suffix == ".xlsx":
        return read_xlsx(path)
    raise ValueError(f"unsupported file type: {suffix or 'no extension'}")
